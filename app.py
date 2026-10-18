import re
from datetime import date, datetime
from io import BytesIO
from zipfile import ZipFile
import unicodedata

from reportlab.lib import colors
from reportlab.lib.enums import TA_CENTER, TA_RIGHT
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import ParagraphStyle
from reportlab.lib.utils import simpleSplit
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont
from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle, KeepTogether
from pathlib import Path

import pandas as pd
import streamlit as st
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

st.set_page_config(page_title="Phần mềm tính tiền lương", page_icon="💰", layout="wide")

LOGO = Path(__file__).resolve().parent / "LOGO.jpg"
if LOGO.exists():
    st.image(str(LOGO), width=180)

st.title("💰 PHẦN MỀM TÍNH TIỀN LƯƠNG")
st.subheader("Song Hào")
st.caption("Quản lý thông tin nhân viên và tính lương cho tối đa 500 người.")
st.info("Đơn vị tiền: triệu đồng/tháng. Công thức thuế và mức giảm trừ được giữ từ mã nguồn gốc; hãy kiểm tra quy định áp dụng trước khi sử dụng chính thức.")

THONG_TIN = ["Họ và tên", "Email cá nhân", "Ngày sinh", "Chức vụ"]
DAU_VAO_TIEN = ["Lương cơ bản", "Tiền thưởng", "Phụ cấp tính BHXH", "Phụ cấp không tính BHXH"]
DAU_VAO = THONG_TIN + DAU_VAO_TIEN + ["Số người phụ thuộc"]
COT_KET_QUA = ["Tổng thu nhập", "Bảo hiểm người lao động đóng", "Giảm trừ bản thân", "Giảm trừ người phụ thuộc", "Thu nhập tính thuế", "Thuế TNCN", "Tiền thực lãnh"]
COT_TIEN = DAU_VAO_TIEN + COT_KET_QUA


def dinh_dang_tien(x):
    return f"{x:,.3f}".replace(",", "X").replace(".", ",").replace("X", ".")


def tinh_thue_tncn(x):
    if x <= 0:
        return 0.0
    if x <= 10:
        return x * 0.05
    if x <= 30:
        return x * 0.10 - 0.5
    if x <= 60:
        return x * 0.20 - 3.5
    if x <= 100:
        return x * 0.30 - 9.5
    return x * 0.35 - 14.5


def tinh_luong(nv):
    tong = sum(float(nv[c]) for c in DAU_VAO_TIEN)
    bh = (float(nv["Lương cơ bản"]) + float(nv["Phụ cấp tính BHXH"])) * 0.105
    giam_tru = 6.2 * int(nv["Số người phụ thuộc"])
    tinh_thue = max(0.0, tong - 15.5 - giam_tru - bh)
    thue = tinh_thue_tncn(tinh_thue)
    return {"Tổng thu nhập": tong, "Bảo hiểm người lao động đóng": bh,
            "Giảm trừ bản thân": 15.5, "Giảm trừ người phụ thuộc": giam_tru,
            "Thu nhập tính thuế": tinh_thue, "Thuế TNCN": thue,
            "Tiền thực lãnh": tong - bh - thue}


def tao_file_excel(df):
    buf = BytesIO()
    with pd.ExcelWriter(buf, engine="openpyxl", date_format="DD/MM/YYYY", datetime_format="DD/MM/YYYY") as writer:
        for sheet_name, startrow in [("Bang luong", 2), ("Chi tiet", 0)]:
            df.to_excel(writer, sheet_name=sheet_name, index=False, startrow=startrow)
            ws = writer.sheets[sheet_name]
            hr = startrow + 1
            if startrow:
                ws["A1"] = "BẢNG TỔNG HỢP LƯƠNG VÀ THUẾ TNCN"
                ws["A1"].font = Font(size=16, bold=True, color="1F4E78")
                ws["A2"] = "Đơn vị: triệu đồng/tháng | Ngày sinh: DD/MM/YYYY"
            ws.row_dimensions[hr].height = 34
            for cell in ws[hr]:
                cell.font = Font(bold=True, color="FFFFFF")
                cell.fill = PatternFill("solid", fgColor="1F4E78")
                cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
            ws.freeze_panes = f"E{hr + 1}"
            ws.auto_filter.ref = f"A{hr}:{get_column_letter(ws.max_column)}{ws.max_row}"
            for idx, name in enumerate(df.columns, start=1):
                ws.column_dimensions[get_column_letter(idx)].width = (
                    28 if name in ["Họ và tên", "Email cá nhân", "Chức vụ"] else
                    17 if name == "Ngày sinh" else 24 if name in COT_TIEN else 20
                )
                for row in range(hr + 1, ws.max_row + 1):
                    cell = ws.cell(row, idx)
                    cell.alignment = Alignment(vertical="center")
                    if name in COT_TIEN:
                        cell.number_format = '#,##0.000;[Red]-#,##0.000;0.000'
                    elif name == "Ngày sinh" and isinstance(cell.value, (date, datetime)):
                        cell.number_format = "DD/MM/YYYY"
            if sheet_name == "Bang luong":
                total_row = ws.max_row + 1
                ws.cell(total_row, 1, "TỔNG CỘNG")
                for idx, name in enumerate(df.columns, start=1):
                    if name in DAU_VAO_TIEN + ["Tổng thu nhập", "Bảo hiểm người lao động đóng", "Thuế TNCN", "Tiền thực lãnh"]:
                        letter = get_column_letter(idx)
                        cell = ws.cell(total_row, idx, f"=SUM({letter}{hr+1}:{letter}{total_row-1})")
                        cell.number_format = '#,##0.000'
                for cell in ws[total_row]:
                    cell.font = Font(bold=True)
                    cell.fill = PatternFill("solid", fgColor="D9EAF7")
    return buf.getvalue()


# PDF: cần cài reportlab và có phông chữ Unicode để hiển thị tiếng Việt.
def dang_ky_font_pdf():
    if "PhieuLuongVN" in pdfmetrics.getRegisteredFontNames():
        return "PhieuLuongVN"
    thu_muc = Path(__file__).resolve().parent
    duong_dan = [
        thu_muc / "NotoSans-Regular.ttf",
        thu_muc / "DejaVuSans.ttf",
        Path("/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf"),
        Path("/usr/share/fonts/truetype/noto/NotoSans-Regular.ttf"),
    ]
    for f in duong_dan:
        if f.is_file():
            pdfmetrics.registerFont(TTFont("PhieuLuongVN", str(f)))
            return "PhieuLuongVN"
    raise FileNotFoundError(
        "Không tìm thấy font Unicode. Hãy đặt file NotoSans-Regular.ttf "
        "hoặc DejaVuSans.ttf cùng thư mục app.py để xuất PDF tiếng Việt."
    )


def dinh_dang_ngay(x):
    return x.strftime("%d/%m/%Y") if isinstance(x, (date, datetime)) else "Chưa cung cấp"


def tao_pdf_nhan_vien(nv):
    from xml.sax.saxutils import escape
    font = dang_ky_font_pdf()
    buf = BytesIO()
    doc = SimpleDocTemplate(buf, pagesize=A4, rightMargin=42, leftMargin=42,
                            topMargin=40, bottomMargin=38, title="Phiếu lương cá nhân")
    normal = ParagraphStyle("NormalVN", fontName=font, fontSize=10, leading=16)
    heading = ParagraphStyle("HeadingVN", parent=normal, fontSize=17,
                             leading=24, alignment=TA_CENTER, textColor=colors.HexColor("#1F4E78"))
    sub = ParagraphStyle("SubVN", parent=normal, fontSize=10,
                         alignment=TA_CENTER, textColor=colors.grey)
    label = ParagraphStyle("LabelVN", parent=normal, fontSize=10)
    bold = ParagraphStyle("BoldVN", parent=normal, fontSize=11, textColor=colors.HexColor("#17365D"))
    def para(value, style=normal):
        return Paragraph(escape(str(value if value is not None else "")), style)
    story = [para("PHIẾU LƯƠNG NHÂN VIÊN", heading),
             para("SONG HÀO | Đơn vị tiền: triệu đồng/tháng", sub), Spacer(1, 18)]
    info = [
        ("Họ và tên", nv["Họ và tên"]),
        ("Email cá nhân", nv.get("Email cá nhân") or "Chưa cung cấp"),
        ("Ngày sinh", dinh_dang_ngay(nv.get("Ngày sinh"))),
        ("Chức vụ", nv.get("Chức vụ") or "Chưa cung cấp"),
        ("Số người phụ thuộc", str(int(nv["Số người phụ thuộc"]))),
    ]
    rows = [[para(k, label), para(v)] for k, v in info]
    tbl = Table(rows, colWidths=[165, 345], hAlign="LEFT")
    tbl.setStyle(TableStyle([("VALIGN", (0,0),(-1,-1),"TOP"),
                             ("BOTTOMPADDING",(0,0),(-1,-1),8),
                             ("LINEBELOW",(0,-1),(-1,-1),0.5,colors.HexColor("#CBD5E1"))]))
    story += [tbl, Spacer(1, 17), para("CHI TIẾT THU NHẬP VÀ KHẤU TRỪ", bold), Spacer(1, 8)]
    fields = ["Lương cơ bản", "Tiền thưởng", "Phụ cấp tính BHXH", "Phụ cấp không tính BHXH",
              "Tổng thu nhập", "Bảo hiểm người lao động đóng", "Giảm trừ bản thân",
              "Giảm trừ người phụ thuộc", "Thu nhập tính thuế", "Thuế TNCN", "Tiền thực lãnh"]
    data = [[para("Khoản mục"), para("Số tiền (triệu đồng)")]]
    for key in fields:
        data.append([para(key), para(dinh_dang_tien(float(nv[key])))])
    money = Table(data, colWidths=[340,170], repeatRows=1)
    money.setStyle(TableStyle([
        ("BACKGROUND",(0,0),(-1,0),colors.HexColor("#1F4E78")),
        ("TEXTCOLOR",(0,0),(-1,0),colors.white),
        ("ROWBACKGROUNDS",(0,1),(-1,-2),[colors.white,colors.HexColor("#F1F5F9")]),
        ("BACKGROUND",(0,-1),(-1,-1),colors.HexColor("#D9EAF7")),
        ("LINEBELOW",(0,-1),(-1,-1),0.8,colors.HexColor("#1F4E78")),
        ("TOPPADDING",(0,0),(-1,-1),8),
        ("BOTTOMPADDING",(0,0),(-1,-1),8),
        ("VALIGN",(0,0),(-1,-1),"MIDDLE"),
    ]))
    # ReportLab Paragraph giữ màu riêng nên tô tiêu đề trực tiếp bằng style chữ trắng.
    data[0] = [para("Khoản mục", ParagraphStyle("WhiteLeft",parent=normal,textColor=colors.white)),
               para("Số tiền (triệu đồng)", ParagraphStyle("WhiteRight",parent=normal,textColor=colors.white))]
    story += [money, Spacer(1, 18),
              para("Phiếu lương được tạo tự động từ dữ liệu tính lương trong ứng dụng.", sub),
              para("Số liệu dựa trên giả định của chương trình, cần kiểm tra trước khi sử dụng chính thức.", sub)]
    doc.build(story)
    return buf.getvalue()


def ten_file_an_toan(name):
    ascii_name = unicodedata.normalize("NFKD", str(name)).encode("ascii", "ignore").decode("ascii")
    return re.sub(r"[^A-Za-z0-9_-]+", "_", ascii_name).strip("_")[:60] or "Nhan_vien"


def tao_zip_phieu_luong(df):
    buf = BytesIO()
    with ZipFile(buf, "w") as zf:
        for i, (_, nv) in enumerate(df.iterrows(), start=1):
            zf.writestr(f"{i:03d}_{ten_file_an_toan(nv['Họ và tên'])}.pdf", tao_pdf_nhan_vien(nv))
    return buf.getvalue()


def tao_du_lieu():
    return pd.DataFrame({
        "Họ và tên": [f"Nhân viên {i+1}" for i in range(500)],
        "Email cá nhân": [""] * 500,
        "Ngày sinh": [""] * 500,
        "Chức vụ": ["Nhân viên"] * 500,
        "Lương cơ bản": [15.0] * 500,
        "Tiền thưởng": [2.0] * 500,
        "Phụ cấp tính BHXH": [1.0] * 500,
        "Phụ cấp không tính BHXH": [0.5] * 500,
        "Số người phụ thuộc": [1] * 500,
    })


if "du_lieu_nhan_vien" not in st.session_state:
    st.session_state.du_lieu_nhan_vien = tao_du_lieu()
if "phien_ban_bang" not in st.session_state:
    st.session_state.phien_ban_bang = 0

st.subheader("1. Chọn số lượng nhân viên")
so_nhan_vien = int(st.number_input("Số nhân viên cần tính", min_value=1, max_value=500, value=10, step=1))
st.subheader("2. Nhập thông tin nhân viên và tiền lương")
st.caption("Chỉnh sửa trực tiếp như Excel; ngày sinh nhập DD/MM/YYYY, ví dụ 09/06/1998. Có thể dán nhiều ô từ Excel.")

bang_goc = st.session_state.du_lieu_nhan_vien.iloc[:so_nhan_vien].copy()
config = {
    "Họ và tên": st.column_config.TextColumn("Họ và tên", required=True, width="medium"),
    "Email cá nhân": st.column_config.TextColumn("Email cá nhân", width="medium"),
    "Ngày sinh": st.column_config.TextColumn("Ngày sinh (DD/MM/YYYY)", width="medium"),
    "Chức vụ": st.column_config.TextColumn("Chức vụ", width="medium"),
    "Số người phụ thuộc": st.column_config.NumberColumn("Số người phụ thuộc", min_value=0, max_value=100, step=1, format="%d"),
}
for c in DAU_VAO_TIEN:
    config[c] = st.column_config.NumberColumn(c, min_value=0.0, step=0.1, format="%.3f")

bang_nhap = st.data_editor(bang_goc, hide_index=True, use_container_width=True,
                          num_rows="fixed", column_config=config,
                          key=f"nhap_{so_nhan_vien}_{st.session_state.phien_ban_bang}")
st.session_state.du_lieu_nhan_vien.iloc[:so_nhan_vien] = bang_nhap.to_numpy()

if st.button("🧮 TÍNH LƯƠNG TOÀN BỘ NHÂN VIÊN", type="primary", use_container_width=True):
   loi = []
    ban_sao = bang_nhap.copy()
    ngay_sinh = []
    for i, nv in ban_sao.iterrows():
        so_dong = i + 1
        ten = str(nv["Họ và tên"]).strip() if pd.notna(nv["Họ và tên"]) else ""
        if not ten:
            loi.append(f"Dòng {so_dong}: thiếu họ tên.")
        email = str(nv["Email cá nhân"]).strip() if pd.notna(nv["Email cá nhân"]) else ""
        if email and not re.fullmatch(r"[^\s@]+@[^\s@]+\.[^\s@]+", email):
            loi.append(f"Dòng {so_dong}: email không hợp lệ.")
        ns = str(nv["Ngày sinh"]).strip() if pd.notna(nv["Ngày sinh"]) else ""
        if ns:
            try:
                d = datetime.strptime(ns, "%d/%m/%Y").date()
                if d > date.today():
                    raise ValueError("future")
                ngay_sinh.append(d)
            except ValueError:
                loi.append(f"Dòng {so_dong}: ngày sinh phải đúng DD/MM/YYYY và không ở tương lai.")
                ngay_sinh.append(None)
        else:
            ngay_sinh.append(None)
        for c in DAU_VAO_TIEN + ["Số người phụ thuộc"]:
            val = pd.to_numeric(nv[c], errors="coerce")
            if pd.isna(val) or val < 0:
                loi.append(f"Dòng {so_dong}: {c} phải là số không âm.")
            elif c == "Số người phụ thuộc" and val != int(val):
                loi.append(f"Dòng {so_dong}: số người phụ thuộc phải là số nguyên.")
    if loi:
        for item in loi[:15]:
            st.error(item)
        if len(loi) > 15:
            st.warning(f"Còn {len(loi)-15} lỗi khác cần sửa.")
    else:
        ban_sao["Ngày sinh"] = ngay_sinh
        for c in DAU_VAO_TIEN:
            ban_sao[c] = pd.to_numeric(ban_sao[c]).astype(float)
        ban_sao["Số người phụ thuộc"] = pd.to_numeric(ban_sao["Số người phụ thuộc"]).astype(int)
        kq = pd.DataFrame([tinh_luong(nv) for _, nv in ban_sao.iterrows()])
        st.session_state.ket_qua_luong = pd.concat([ban_sao.reset_index(drop=True), kq], axis=1)
        st.session_state.pop("zip_phieu_luong", None)
        st.success("Đã tính xong bảng lương. Bạn có thể tải file Excel bên dưới.")

if "ket_qua_luong" in st.session_state:
    df = st.session_state.ket_qua_luong
    st.divider()
    st.subheader("3. Kết quả tính lương")
    a, b, c, d = st.columns(4)
    a.metric("Số nhân viên", len(df))
    b.metric("Tổng thu nhập", f'{dinh_dang_tien(df["Tổng thu nhập"].sum())} triệu')
    c.metric("Tổng thuế TNCN", f'{dinh_dang_tien(df["Thuế TNCN"].sum())} triệu')
    d.metric("Tổng thực lãnh", f'{dinh_dang_tien(df["Tiền thực lãnh"].sum())} triệu')
    hien_thi = df.copy()
    hien_thi["Ngày sinh"] = hien_thi["Ngày sinh"].apply(lambda x: x.strftime("%d/%m/%Y") if isinstance(x, (date, datetime)) else "")
    st.dataframe(hien_thi, use_container_width=True, hide_index=True,
                 column_config={c: st.column_config.NumberColumn(c, format="%.3f") for c in COT_TIEN})
    st.subheader("4. Xuất dữ liệu")
    st.download_button("📥 TẢI FILE EXCEL (.xlsx)", data=tao_file_excel(df),
                       file_name="bang_luong_nhan_vien.xlsx",
                       mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                       use_container_width=True)
    st.download_button("📥 TẢI FILE CSV", data=hien_thi.to_csv(index=False).encode("utf-8-sig"),
                       file_name="bang_luong_nhan_vien.csv", mime="text/csv", use_container_width=True)
    st.subheader("5. Xuất phiếu lương PDF từng nhân viên")
    chon_nv = st.selectbox("Chọn nhân viên để tải phiếu lương PDF", options=list(range(len(df))),
                           format_func=lambda i: f"{i+1:03d} - {df.iloc[i]['Họ và tên']}")
    try:
        pdf = tao_pdf_nhan_vien(df.iloc[chon_nv])
        st.download_button("📄 TẢI PHIẾU LƯƠNG PDF CỦA NHÂN VIÊN ĐÃ CHỌN",
                           data=pdf, file_name=f"phieu_luong_{chon_nv+1:03d}_{ten_file_an_toan(df.iloc[chon_nv]['Họ và tên'])}.pdf",
                           mime="application/pdf", use_container_width=True)
        if st.button("📦 TẠO FILE ZIP CHỨA PDF CỦA TẤT CẢ NHÂN VIÊN", use_container_width=True):
            with st.spinner("Đang tạo các phiếu lương PDF..."):
                st.session_state.zip_phieu_luong = tao_zip_phieu_luong(df)
        if "zip_phieu_luong" in st.session_state:
            st.download_button("⬇️ TẢI FILE ZIP PHIẾU LƯƠNG", data=st.session_state.zip_phieu_luong,
                               file_name="phieu_luong_tat_ca_nhan_vien.zip", mime="application/zip",
                               use_container_width=True)
    except (FileNotFoundError, ImportError) as exc:
        st.error(f"Chưa thể tạo PDF: {exc}")
    st.caption("Sau khi sửa dữ liệu nhập, nhấn TÍNH LƯƠNG lại để cập nhật kết quả và file xuất.")

with st.expander("📘 Công thức đang sử dụng"):
    st.markdown("""
    - **Tổng thu nhập:** Lương cơ bản + Thưởng + Phụ cấp tính BHXH + Phụ cấp không tính BHXH.
    - **Bảo hiểm người lao động đóng:** (Lương cơ bản + Phụ cấp tính BHXH) × 10,5%.
    - **Thu nhập tính thuế:** MAX(0; Tổng thu nhập − 15,5 − 6,2 × Số người phụ thuộc − Bảo hiểm).
    - **Thuế TNCN:** Tính theo biểu thuế lũy tiến được giữ từ chương trình gốc.
    - **Tiền thực lãnh:** Tổng thu nhập − Bảo hiểm − Thuế TNCN.
      """)
